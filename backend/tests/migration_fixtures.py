"""Synthetic copies of the migration inputs, shaped like the real ones."""
import hashlib
import json
import sqlite3
from datetime import datetime
from pathlib import Path

from tests.pdf_fixture import make_pdf
from tutoring.passages import split_sentences

GENGHIS_PASSAGE = ("Genghis Khan, Emperor of All Men\n\nThe Mongols came from the wild, cold mountains north of China. "
                   "They lived in felt tents.\n\nOne man spoke up. He knocked the Mongol off his horse.")


def textbook_pdf() -> bytes:
    # Book page n is PDF page n + 2 (two front-matter pages), like SOTW's offset of 15.
    pages = [["Front matter"], ["Contents"],
             ["Robin Hood", "Robin lived in the royal forests."],
             ["The Scattering of the Jews", "The Romans scattered the Jews."],
             ["A Tale of the Diaspora", "The rabbi ate the paper."],
             ["Genghis Khan, Emperor of All Men", "The Mongols came from the wild, cold mountains north of China.",
              "They lived in felt tents."],
             ["One man spoke up. He knocked the Mongol off his horse.", "The Mongol Conquest of China",
              "Kublai Khan became emperor."]]
    return make_pdf(pages)


def build_history_folder(root: Path, voice="en-US-Neural2-J", rate=0.9) -> dict:
    import openpyxl
    root.mkdir(parents=True, exist_ok=True)
    book = textbook_pdf()
    (root / "Story of the World V.2").write_bytes(book)
    tests = make_pdf([["Chapter 21 test", "1. Who was Genghis Khan?"]] * 6)
    (root / "SOTW2_Tests.pdf").write_bytes(tests)
    config = {"subject": "History", "learner": "Lucas", "tracker": "outputs/01a07bde/Lucas_History_Tracker.xlsx",
              "curriculum": "subject/curriculum.json", "textbook": "Story of the World V.2.pdf",
              "tests": "SOTW2_Tests.pdf", "answer_key": "teacher_resources/sotw2_test_answer_key.sqlite",
              "cache": "teacher_resources/textbook", "tts": {"project": "gen-lang-client-0088393168", "voice": voice,
                                                            "rate": rate, "monthly_limit": 150000}}
    (root / "tutor.json").write_text(json.dumps(config))
    topics = [("C19-T03", 19, "Robin Hood / Robin Hood and the Butcher", 1, 1, "The Scattering of the Jews"),
              ("C20-T01", 20, "The Scattering of the Jews", 2, 2, "A Tale of the Diaspora"),
              ("C20-T02", 20, "A Tale of the Diaspora / The Clever Rabbi of Cordova", 3, 3,
               "Genghis Khan, Emperor of All Men"),
              ("C21-T01", 21, "Genghis Khan, Emperor of All Men", 4, 5, "The Mongol Conquest of China"),
              ("C21-T02", 21, "The Mongol Conquest of China", 5, 5, None)]
    curriculum = {
        "chapters": [{"chapter": c, "title": f"Chapter {c}", "book_start": 1, "book_end": 5, "test_pdf_start": 2,
                      "test_pdf_end": 4, "test_print_start": 1, "test_print_end": 3, "key_pdf": [6]}
                     for c in (19, 20, 21)],
        "topics": [{"id": key, "chapter": chapter, "title": title, "start": start, "end": end, "pdf_start": start + 2,
                    "pdf_end": end + 2, "order": index + 1, **({"stop_before": stop} if stop else {})}
                   for index, (key, chapter, title, start, end, stop) in enumerate(topics)],
        "sources": {"text": "Story of the World V.2.pdf", "tests": "SOTW2_Tests.pdf"},
    }
    (root / "subject").mkdir()
    (root / "subject" / "curriculum.json").write_text(json.dumps(curriculum))
    (root / "subject" / "key_exceptions.json").write_text(json.dumps([
        {"chapter": 19, "question": 4, "issue": "The key says Sherwood Forest.",
         "action": "Grade against the passage.", "source_reference": "key p. 173"}]))
    (root / "subject" / "question-mappings").mkdir()
    (root / "subject" / "question-mappings" / "C21.json").write_text(json.dumps({
        "topics": {"C21-T01": [1, 3], "C21-T02": [4]},
        "evidence": {"1": "They lived in felt tents.", "3": "One man spoke up.", "4": "Kublai Khan became emperor."},
        "question_notes": {"4": "Accept a supported explanation."}}))
    (root / "subject" / "reviews").mkdir()
    (root / "subject" / "reviews" / "C21.json").write_text(json.dumps({
        "questions": {"1": ["C21-T01", "felt tents"], "3": ["C21-T01", "spoke up"], "4": ["C21-T02", "Kublai"]}}))

    tr = root / "teacher_resources"
    (tr / "textbook" / "reviewed").mkdir(parents=True)
    key = sqlite3.connect(tr / "sotw2_test_answer_key.sqlite")
    key.executescript("""
        create table chapters(chapter integer, key_pdf_page integer, key_printed_page integer, answer_count integer);
        create table answers(chapter integer, question_start integer, question_end integer, answer text);
        create table exceptions(chapter integer, question integer, issue text, action text, source_reference text);
        insert into chapters values (19, 6, 173, 17), (21, 6, 173, 18);
        insert into answers values (19, 4, 4, 'Sherwood Forest'), (21, 1, 1, 'SECRET: a Mongol leader'),
                                   (21, 3, 3, 'He knocked the Mongol off his horse'), (21, 4, 4, 'Kublai Khan');
    """)
    key.commit()
    key.close()
    book_sha = hashlib.sha256(book).hexdigest()
    overlay = {"schema_version": 2, "topic_id": "C21-T01", "chapter": 21, "title": "Genghis Khan, Emperor of All Men",
               "printed_pages": [4, 5], "pdf_pages": [6, 7], "start_at": "Genghis Khan, Emperor of All Men",
               "stop_before": "The Mongol Conquest of China", "passage": GENGHIS_PASSAGE,
               "passage_sha256": hashlib.sha256(GENGHIS_PASSAGE.encode()).hexdigest(),
               "sentences": split_sentences(GENGHIS_PASSAGE),
               "selected_images": [{"asset": "images/p6.png", "pdf_page": 6, "classification": "illustration",
                                    "crop": None, "alt": "Genghis Khan on horseback", "caption": "(book p. 4)",
                                    "before_sentence": 2}],
               "validation": {"status": "verified", "issues": []},
               "provenance": {"pdf": book_sha, "extraction": {"adapter": "sotw2", "version": 2}},
               "review": {"reviewed_at": "2026-09-24T21:12:52+00:00", "pdf_pages": [6, 7],
                          "notes": "Visually checked all assigned pages."}}
    (tr / "textbook" / "reviewed" / "C21-T01.json").write_text(json.dumps(overlay))

    reading = root / "assets" / "reading"
    reading.mkdir(parents=True)
    rate_text = "0.9"
    legacy_hash = hashlib.sha256(f"{voice}\n{rate_text}\n{GENGHIS_PASSAGE}".encode()).hexdigest()
    count = len(overlay["sentences"])
    (reading / "narration-5e9d-google-neural2-manifest.json").write_text(json.dumps({
        "version": 2, "label": "c21-t01", "provider": "Google Cloud Text-to-Speech", "voice": voice,
        "speakingRate": rate, "characterCount": len(GENGHIS_PASSAGE), "sentenceCount": count,
        "passageHash": legacy_hash, "tracks": [{"src": "assets/reading/narration-5e9d-part-1.mp3", "startSentence": 0,
                                                "sentenceStarts": [i * 3.0 for i in range(count)]}]}))
    (reading / "c20-t02-clever-rabbi-google-neural2-manifest.json").write_text(json.dumps({
        "version": 2, "label": "c20-t02-clever-rabbi", "provider": "Google Cloud Text-to-Speech", "voice": voice,
        "speakingRate": rate, "sentenceCount": 3, "passageHash": "0" * 64,
        "tracks": [{"src": "assets/reading/c20-t02-part-1.mp3", "startSentence": 0, "sentenceStarts": [0, 1, 2]}]}))
    for name in ("narration-5e9d-part-1.mp3", "c20-t02-part-1.mp3", "robin-hood-elevenlabs-part-1.mp3",
                 "robin-hood-elevenlabs-part-2.mp3", "gemini-preview-1.wav"):
        (reading / name).write_bytes(b"ID3" + name.encode())

    (root / "context").mkdir()
    (root / "context" / "recovered-genghis.json").write_text(json.dumps({
        "status": "unfinished", "topic_id": "C21-T01", "phase": "discussion", "source_thread": "01a07c05",
        "basis": "Exact learner responses recovered from the parent-approved source conversation.",
        "reading_signal": "im ready", "next_prompt": "What did the man who spoke up do next?", "minutes": None,
        "observations": [
            {"prompt": "What did he do with the tribes?", "first_response": "he made all of them join together and attacked china"},
            {"prompt": "Why were the united tribes harder to defeat?", "first_response": "yes absulootly",
             "help": "What could the tribes do together?", "revision": "they coudl oncquor way more",
             "explanation_supplied": "Tutor explained larger army."}]}))
    (root / "context" / "session.json").write_text(json.dumps({"status": "saved", "session_id": "2026-09-23-01",
                                                               "topic_id": "C20-T02"}))
    (root / "evidence" / "2026-09-11").mkdir(parents=True)

    workbook = openpyxl.Workbook()
    start = workbook.active
    start.title = "Start"
    for row in (["Start Lucas's history learning record"], ["Learner", "Lucas"], ["Age at setup", 9],
                ["Tutoring language", "English"], ["Sessions per week", 3], ["Parent present", "Yes"],
                ["Recorded progress", "Count"], ["Topics completed", 3]):
        start.append(row)
    sheet = workbook.create_sheet("Topics")
    sheet.append(["Topics Topic plan and understanding"])
    sheet.append(["Topic ID", "Chapter", "Reading topic", "Book pages", "PDF pages", "Stop before", "Coverage",
                  "Understanding", "Date completed", "Last checked", "Evidence / session ID", "Next teaching step",
                  "Reading order"])
    rows = {"C19-T03": ("Completed", "Needs help", datetime(2026, 9, 11), "2026-09-11-01"),
            "C20-T01": ("Completed", "Needs help", datetime(2026, 9, 14), "2026-09-14-01"),
            "C20-T02": ("Completed", "Independent", datetime(2026, 9, 24), "2026-09-23-01")}
    for index, (key, *_rest) in enumerate(topics):
        coverage, understanding, day, session = rows.get(key, ("Planned", "Not assessed", None, None))
        sheet.append([key, _rest[0], _rest[1], "", "", "", coverage, understanding, day, day, session, None, index + 1])
    sheet = workbook.create_sheet("Chapters")
    sheet.append(["Chapters Chapter tests"])
    sheet.append(["Chapter", "Chapter title", "Book pages", "Test PDF pages", "Test book pages", "Key PDF page",
                  "Test status", "Test date", "Session ID", "Items to revisit", "Parent / tutor note"])
    sheet.append([19, "A New Kind of King", "", "", "", 6, "Completed", datetime(2026, 9, 11), "2026-09-11-01",
                  "Q4 forest", "16 correct"])
    sheet = workbook.create_sheet("Sessions")
    sheet.append(["Sessions Lesson history"])
    sheet.append(["Session ID", "Date", "Topic ID", "Session type", "Reading mode", "Minutes",
                  "What Lucas explained / did", "What was difficult", "Help that worked", "Next lesson / action",
                  "Parent observation", "Photo / note path"])
    sheet.append(["2026-09-11-01", datetime(2026, 9, 11), "C19-T03", "Chapter test", None, None,
                  "Completed all 17 items", "Q4 forest wording", "Passage pointer", "Begin C20-T01", "Three photos",
                  "evidence/2026-09-11/page-71.jpg"])
    sheet.append(["2026-09-23-01", datetime(2026, 9, 23), "C20-T02", "Reading", None, None,
                  "Retold the Clever Rabbi", "Yohanan's school", "Chronology", "Begin Chapter 21", None, None])
    sheet = workbook.create_sheet("Answers")
    sheet.append(["Answers Handwritten answers and corrections"])
    sheet.append(["Answer ID", "Session ID", "Topic ID", "Chapter", "Question / page", "First answer (verbatim)",
                  "Reading certainty", "First result", "Hint given", "Corrected answer", "After-help result",
                  "Recheck date", "Photo path", "Source / key / tutor note"])
    sheet.append(["2026-09-11-01-Q01", "2026-09-11-01", "C19-T03", 19, "Q1 / book 71 (PDF 74)", "Lionhearted",
                  "Clear", "Correct", None, None, None, None, "evidence/2026-09-11/page-71.jpg", "Key agrees"])
    sheet.append(["2026-09-11-01-Q04", "2026-09-11-01", "C19-T03", 19, "Q4 / book 71 (PDF 74)",
                  "The forest of England", "Clear", "Partly correct", "Look again at Q4", "Forest of Nottingham",
                  "With help", datetime(2026, 9, 12), "evidence/2026-09-11/page-71.jpg", "Key conflict"])
    sheet = workbook.create_sheet("Review")
    sheet.append(["Review Review queue"])
    sheet.append(["Review ID", "Topic ID", "Concept to revisit", "Observed difficulty / evidence", "Last practice",
                  "Interval (days)", "Next due", "Status", "Last outcome", "Next prompt / strategy", "Action",
                  "Evidence / session ID"])
    sheet.append(["2026-09-14-01-R01", "C20-T01", "Yohanan ben Zakkai's work", "Confused school with synagogues",
                  datetime(2026, 9, 24), 2, datetime(2026, 9, 26), "Open", "With help", "What did Yohanan do?",
                  "Scheduled", "2026-09-23-01"])
    sheet.append([None] * 12)
    tracker = root / config["tracker"]
    tracker.parent.mkdir(parents=True)
    workbook.save(tracker)
    drive_paths = {"evidence/2026-09-11/page-71.jpg": "photoFileId0001",
                   "assets/reading/narration-5e9d-part-1.mp3": "audioFileId0001",
                   "assets/reading/c20-t02-part-1.mp3": "audioFileId0002",
                   "assets/reading/robin-hood-elevenlabs-part-1.mp3": "audioFileId0003",
                   "assets/reading/robin-hood-elevenlabs-part-2.mp3": "audioFileId0004"}
    return {"book": book, "tests": tests, "book_sha": book_sha, "drive_paths": drive_paths, "tracker": tracker}


def build_hosted_export(root: Path, book: bytes) -> Path:
    """A small export shaped like migration.export_hosted output."""
    tables = {
        "children": [{"id": 1, "owner_id": "u", "name": "Lucas", "nickname": None, "color": "#4A90D9",
                      "grade_year": "4th", "created_at": "2026-08-09T20:00:00+00:00"},
                     {"id": 2, "owner_id": "u", "name": "Mila", "nickname": None, "color": "#E88AB5",
                      "grade_year": "2nd", "created_at": "2026-08-09T20:00:00+00:00"}],
        "subjects": [{"id": 5, "child_id": 1, "name": "History", "weight": 1.0, "slot_type": "A", "end_date": None,
                      "created_at": "2026-08-09T20:00:00+00:00"},
                     {"id": 6, "child_id": 2, "name": "Language Arts", "weight": 1.0, "slot_type": "B",
                      "end_date": None, "created_at": "2026-08-09T20:00:00+00:00"}],
        "documents": [{"id": 3, "owner_id": "u", "blob_path": "documents/abc.pdf",
                       "original_filename": "Story of the World V.2.pdf", "size_bytes": len(book), "page_count": 7,
                       "sha256": hashlib.sha256(book).hexdigest(), "status": "ready",
                       "created_at": "2026-08-09T20:00:00+00:00"}],
        "curriculum_topics": [{"id": 40, "subject_id": 5, "title": "Genghis Khan, Emperor of All Men", "page_start": 4,
                               "page_end": 5, "complexity": 1, "completed": False, "completed_at": None,
                               "language": "en", "chapter_order": 54, "pdf_filename": "Story of the World V.2.pdf",
                               "document_id": 3, "pdf_page_offset": 2, "is_core": True,
                               "created_at": "2026-08-09T20:00:00+00:00"},
                              {"id": 41, "subject_id": 6, "title": "Lesson 1", "page_start": 1, "page_end": 2,
                               "complexity": 1, "completed": True, "completed_at": "2026-08-20T13:00:00+00:00",
                               "language": "en", "chapter_order": 1, "pdf_filename": None, "document_id": None,
                               "pdf_page_offset": 0, "is_core": True, "created_at": "2026-08-09T20:00:00+00:00"}],
        "time_windows": [{"id": 1, "child_id": 1, "weekday": 0, "start_time": "09:00", "end_time": "10:00"}],
        "blocked_days": [{"id": 1, "owner_id": "u", "child_id": None, "date": "2026-10-12", "block_type": "holiday",
                          "note": "Feriado", "created_at": "2026-08-09T20:00:00+00:00"}],
        "scheduled_slots": [{"id": 9, "child_id": 2, "subject_id": 6, "topic_id": 41, "date": "2026-08-20",
                             "time_start": "09:00", "time_end": "09:30", "page_from": 1, "page_to": 2}],
        "completions": [{"id": 4, "slot_id": 9, "completed_at": "2026-08-20T13:00:00+00:00"}],
        "canvas_inserts": [],
        "canvas_ai_content": [{"id": 2, "topic_id": 41, "page_start": 1, "page_end": 2, "content_type": "quiz",
                               "content": "[{\"question\": \"Q?\"}]", "created_at": "2026-08-21T10:00:00+00:00"}],
        "app_settings": [{"id": 1, "owner_id": "u", "key": "SCHOOL_YEAR_END", "value": "2026-12-18",
                          "updated_at": "2026-08-10T10:00:00+00:00"}],
        "pdf_page_annotations": [{"id": 7, "owner_id": "u", "child_id": 1, "document_id": 3, "page_number": 6,
                                  "strokes": [{"id": "5b2c", "color": "#111827", "width": 0.004,
                                               "points": [[0.1, 0.2], [0.3, 0.4]]}],
                                  "revision": 4, "created_at": "2026-08-10T15:00:00+00:00",
                                  "updated_at": "2026-08-11T15:00:00+00:00"}],
    }
    (root / "tables").mkdir(parents=True)
    manifest = {"tables": {}, "books": {"3": {"file": f"{hashlib.sha256(book).hexdigest()}.pdf",
                                              "sha256": hashlib.sha256(book).hexdigest(), "size": len(book),
                                              "matches_record": True}}}
    for name, rows in tables.items():
        (root / "tables" / f"{name}.json").write_text(json.dumps(rows))
        manifest["tables"][name] = {"count": len(rows), "ids": [r["id"] for r in rows]}
    (root / "books").mkdir()
    (root / "books" / f"{hashlib.sha256(book).hexdigest()}.pdf").write_bytes(book)
    (root / "manifest.json").write_text(json.dumps(manifest))
    return root
